"""Build a course chatbot's knowledge file from verified course materials.

Generalized from a real course-chatbot build (2026-10).
Standard: 00-System/COURSE-COMPANION-STANDARD.md · Skill: .claude/skills/course-companion/

Inputs (all instructor-authored or page-verified; never answer keys, exam items or licensed full texts):
  --decks  GLOB      .pptx decks; slide text only (no speaker notes). A "_v3" deck replaces its "_v2" twin.
  --guide  FILE.md   student study guides (repeatable); split into sections at "#"/"##" headings.
  --weekmap JSON     optional {"09":"w09", ...}: two-digit week number in the deck filename -> page week id.
  --out    FILE      output JSON: {"version", "items":[{id,type,week,title,cite,text,kw}]}

Example:
  python3 00-System/Scripts/build_companion_kp.py \
    --decks "03-Projects/X/Presentations/Week-*_v2.pptx" \
    --guide "03-Projects/X/Exams/Midterm-Guide-Students.md" \
    --out   "03-Projects/X/suite/data/companion-kp.json"
"""
import argparse, glob, json, os, re, datetime
from pptx import Presentation

ap = argparse.ArgumentParser()
ap.add_argument("--decks", default="")
ap.add_argument("--guide", action="append", default=[])
ap.add_argument("--weekmap", default="")
ap.add_argument("--out", required=True)
a = ap.parse_args()
weekmap = json.loads(a.weekmap) if a.weekmap else {}

SKIP = re.compile(r"^(SOURCES|REFERENCES|LOOKING AHEAD)", re.I)
srcline = lambda t: " | " in t or bool(re.search(r"\bpp?\. ?\d|doi|^Source|\(\d{4}[a-z]?(, [A-Z][a-z]+\.? \d+)?\)\. |\d+ F\. ?(Supp|2d|3d)|\b[A-Z][a-z]{2,8}\.? \d{1,2}, (19|20)\d\d", t))
items = []

decks = sorted(glob.glob(a.decks)) if a.decks else []
for f in list(decks):
    v3 = f.replace("_v2.pptx", "_v3.pptx")
    if v3 != f and os.path.exists(v3): decks[decks.index(f)] = v3
for f in decks:
    base = os.path.basename(f)
    m = re.match(r"Week-(\d\d)", base); wk = m.group(1) if m else "00"
    wid = weekmap.get(wk, "w" + wk)
    kind = "article analysis" if "ARTICLE-ANALYSIS" in base.upper() else "lecture"
    label = f"Week {int(wk)} {kind}"
    for i, s in enumerate(Presentation(f).slides, 1):
        if i == 1: continue                                   # title slide
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                t = " ".join(p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip())
                if t and not re.fullmatch(r"\d+", t) and not re.match(r"^[A-Z]{2,5} \d{3}.{0,3}·", t): texts.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows: texts.append(" | ".join(c.text.strip() for c in r.cells))
        if not texts: continue
        eyebrow = next((t for t in texts if t.isupper() and len(t) < 80), "")
        if SKIP.match(eyebrow): continue
        title = next((t for t in texts if t != eyebrow and not t.isupper() and 12 < len(t) < 160 and not srcline(t)), next((t for t in texts if t != eyebrow), texts[0]))
        body = re.sub(r"\s+", " ", " · ".join(t for t in texts if t not in (eyebrow, title))).strip()
        items.append({"id": f"s-{wk}-{'A' if kind[0] == 'a' else 'L'}-{i}", "type": "slides", "week": wid,
                      "title": title, "cite": f"{label}, slide {i}", "text": body[:1400],
                      "kw": (eyebrow + " " + title).lower()})

for gpath in a.guide:
    g = re.sub(r"^---.*?---\s*", "", open(gpath, encoding="utf8").read(), flags=re.S)
    gname = os.path.splitext(os.path.basename(gpath))[0].replace("-", " ")
    sec, sub, buf = "", "", []
    def flush():
        if buf and "".join(buf).strip():
            txt = re.sub(r"\s+", " ", re.sub(r"[*|#]|-{3,}", " ", "\n".join(buf))).strip()
            items.append({"id": f"g-{len(items)}", "type": "study", "week": "", "title": (sub or sec).strip(),
                          "cite": gname + " · " + (sec.split(". ", 1)[-1] if sec else ""), "text": txt[:1600],
                          "kw": (gname + " study guide " + sec + " " + sub).lower()})
    for line in g.splitlines():
        if line.startswith("# "): flush(); buf = []; sec, sub = line[2:], ""
        elif line.startswith("## "): flush(); buf = []; sub = line[3:]
        else: buf.append(line)
    flush()

os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
json.dump({"version": datetime.date.today().isoformat(), "items": items}, open(a.out, "w", encoding="utf8"), ensure_ascii=False)
print(len(items), "items ->", a.out, f"({os.path.getsize(a.out)//1024} KB)")
